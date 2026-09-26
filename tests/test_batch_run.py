import csv
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "verifying-sales-leads" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import leadlib  # noqa: E402
import manage_batches  # noqa: E402


class BatchRunTests(unittest.TestCase):
    def campaign(self):
        return {
            "id": "large-fixture",
            "research_mode": "batch",
            "product": {"name": "抹茶", "keywords": ["matcha"]},
            "market": {"name": "アメリカ", "keywords": ["United States"]},
            "objective": "find_importers",
            "buyer_types": [{"name": "輸入業者", "keywords": ["importer"]}],
            "excluded_types": [],
            "target_accepted": 2,
            "max_candidates": 4,
            "batch_size": 2,
            "max_pages_per_company": 5,
        }

    def test_campaign_limits_depend_on_mode(self):
        standard_at_limit = self.campaign()
        standard_at_limit["research_mode"] = "standard"
        standard_at_limit["target_accepted"] = 50
        standard_at_limit["max_candidates"] = 100
        self.assertEqual(leadlib.validate_campaign(standard_at_limit), [])

        standard = self.campaign()
        standard["research_mode"] = "standard"
        standard["target_accepted"] = 51
        self.assertIn("target_accepted must be an integer from 1 to 50", leadlib.validate_campaign(standard))

        large = self.campaign()
        large["target_accepted"] = 500
        large["max_candidates"] = 2500
        large["batch_size"] = 50
        self.assertEqual(leadlib.validate_campaign(large), [])

    def test_large_run_deduplicates_resumes_and_exports(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            campaign_path = root / "source-campaign.json"
            campaign_path.write_text(json.dumps(self.campaign()), encoding="utf-8")
            run_dir = root / "run"
            initial = manage_batches.initialize(campaign_path, run_dir)
            self.assertEqual(initial["next_action"], "discover_more_candidates")

            candidates_path = root / "discovered.json"
            candidates_path.write_text(
                json.dumps(
                    [
                        {"url": "https://example1.com/", "discovery_query": "matcha importer"},
                        {"url": "https://www.example1.com/about", "discovery_query": "duplicate"},
                        {"url": "https://example2.com/", "discovery_query": "matcha wholesale"},
                    ]
                ),
                encoding="utf-8",
            )
            added = manage_batches.add_candidates(run_dir, candidates_path)
            self.assertEqual(added["added"], 2)
            self.assertEqual(added["duplicates_skipped"], 1)

            batch = manage_batches.next_batch(run_dir)
            self.assertEqual(batch["status"], "batch_created")
            self.assertEqual(batch["candidate_count"], 2)
            resumed = manage_batches.next_batch(run_dir)
            self.assertEqual(resumed["status"], "resume_open_batch")
            self.assertEqual(resumed["batch_id"], batch["batch_id"])

            selected = json.loads(Path(batch["candidates_file"]).read_text(encoding="utf-8"))
            crawls = []
            assessments = []
            for index, candidate in enumerate(selected, start=1):
                page_url = candidate["url"] + "about"
                text = "Japanese matcha. Wholesale importer. Based in the United States."
                crawls.append(
                    {
                        "candidate": candidate,
                        "canonical_url": candidate["url"],
                        "final_url": page_url,
                        "site_status": "active",
                        "pages": [{"url": page_url, "text": text}],
                        "emails": [],
                        "contact_forms": [],
                        "errors": [],
                        "checked_at": "2026-09-25T00:00:00+00:00",
                    }
                )
                assessments.append(
                    {
                        "candidate_url": candidate["url"],
                        "company_name": f"Example {index}",
                        "country": "アメリカ",
                        "company_type": "輸入卸",
                        "company_overview_ja": "抹茶の輸入卸。",
                        "recommendation": "accepted",
                        "claims": [
                            {"type": "product", "evidence_url": page_url, "evidence_text_original": "Japanese matcha.", "evidence_text_ja": "日本産抹茶。"},
                            {"type": "buyer_role", "evidence_url": page_url, "evidence_text_original": "Wholesale importer.", "evidence_text_ja": "輸入卸。"},
                            {"type": "target_market", "evidence_url": page_url, "evidence_text_original": "Based in the United States.", "evidence_text_ja": "米国拠点。"},
                        ],
                        "uncertainties": [],
                        "rejection_reason": None,
                    }
                )
            crawls_path = root / "crawls.json"
            assessments_path = root / "assessments.json"
            crawls_path.write_text(json.dumps(crawls), encoding="utf-8")
            assessments_path.write_text(json.dumps(assessments), encoding="utf-8")
            summary = manage_batches.complete_batch(
                run_dir, batch["batch_id"], crawls_path, assessments_path
            )
            self.assertTrue(summary["target_met"])
            self.assertEqual(summary["candidates_checked"], 2)
            self.assertEqual(summary["completed_batches"], 1)
            self.assertEqual(manage_batches.next_batch(run_dir)["stop_reason"], "target_met")

            with (run_dir / "leads.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(len(rows), 2)
            self.assertEqual(len(json.loads((run_dir / "audit.json").read_text(encoding="utf-8"))), 2)
            self.assertTrue((run_dir / "leads.xlsx").exists())
            from openpyxl import load_workbook

            workbook = load_workbook(run_dir / "leads.xlsx", read_only=True)
            self.assertEqual(workbook.sheetnames, ["営業リスト"])
            headers = [cell.value for cell in next(workbook["営業リスト"].iter_rows(min_row=4, max_row=4))]
            self.assertNotIn("人の判定", headers)


if __name__ == "__main__":
    unittest.main()
