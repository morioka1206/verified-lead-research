#!/usr/bin/env python3
"""Create clean sales and optional quality-review Excel workbooks."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


SALES_HEADERS = [
    "No.", "会社名", "公式URL", "会社種別", "会社概要", "メール",
    "問い合わせフォームURL", "商材との関係", "対象地域", "確認日時",
]

REVIEW_HEADERS = [
    "No.", "人の判定", "確認メモ", "会社名", "公式URL", "会社種別", "会社概要",
    "メール", "問い合わせフォームURL", "商材の証拠", "商材の証拠URL",
    "会社の役割の証拠", "会社の役割の証拠URL", "対象地域の証拠",
    "対象地域の証拠URL", "確認日時",
]


def claim_text(record: dict[str, Any], claim_type: str) -> str:
    return " / ".join(
        claim.get("evidence_text_ja") or claim.get("evidence_text_original", "")
        for claim in record.get("claims") or []
        if claim.get("type") == claim_type
    )


def claim_urls(record: dict[str, Any], claim_type: str) -> str:
    return " / ".join(
        dict.fromkeys(
            claim.get("evidence_url", "")
            for claim in record.get("claims") or []
            if claim.get("type") == claim_type and claim.get("evidence_url")
        )
    )


def _emails(record: dict[str, Any]) -> str:
    return " / ".join(item["email"] for item in record.get("emails") or [])


def _contact_forms(record: dict[str, Any]) -> str:
    return " / ".join(item["url"] for item in record.get("contact_forms") or [])


def sales_row(record: dict[str, Any], number: int) -> list[str | int]:
    return [
        number,
        record.get("company_name") or "",
        record.get("canonical_url") or "",
        record.get("company_type") or "",
        record.get("company_overview_ja") or "",
        _emails(record),
        _contact_forms(record),
        claim_text(record, "product"),
        claim_text(record, "target_market"),
        record.get("checked_at") or "",
    ]


def review_row(
    record: dict[str, Any],
    number: int,
    human_labels: dict[str, tuple[str, str]] | None = None,
) -> list[str | int]:
    verdict, notes = (human_labels or {}).get(record.get("canonical_url") or "", ("", ""))
    return [
        number,
        verdict,
        notes,
        record.get("company_name") or "",
        record.get("canonical_url") or "",
        record.get("company_type") or "",
        record.get("company_overview_ja") or "",
        _emails(record),
        _contact_forms(record),
        claim_text(record, "product"),
        claim_urls(record, "product"),
        claim_text(record, "buyer_role"),
        claim_urls(record, "buyer_role"),
        claim_text(record, "target_market"),
        claim_urls(record, "target_market"),
        record.get("checked_at") or "",
    ]


def _stable_key(record: dict[str, Any]) -> str:
    identity = record.get("canonical_url") or record.get("company_name") or ""
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


def select_quality_sample(records: list[dict[str, Any]], sample_size: int = 10) -> list[dict[str, Any]]:
    """Select a deterministic mix of difficult and representative accepted leads."""
    if sample_size < 1:
        raise ValueError("sample_size must be at least 1")
    accepted = [record for record in records if record.get("verification_status") == "accepted"]
    if len(accepted) <= sample_size:
        return accepted

    difficult = sorted(
        (record for record in accepted if record.get("uncertainties")),
        key=lambda record: (-len(record.get("uncertainties") or []), _stable_key(record)),
    )
    difficult_limit = min(len(difficult), (sample_size + 1) // 2)
    selected = difficult[:difficult_limit]
    selected_urls = {record.get("canonical_url") for record in selected}

    groups: dict[str, list[dict[str, Any]]] = {}
    for record in accepted:
        if record.get("canonical_url") in selected_urls:
            continue
        groups.setdefault(record.get("company_type") or "未分類", []).append(record)
    for group in groups.values():
        group.sort(key=_stable_key)

    group_names = sorted(groups)
    while len(selected) < sample_size and group_names:
        next_group_names = []
        for name in group_names:
            if groups[name] and len(selected) < sample_size:
                selected.append(groups[name].pop(0))
            if groups[name]:
                next_group_names.append(name)
        group_names = next_group_names
    return selected


def _add_table(sheet, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def write_sales_workbook(records: list[dict[str, Any]], output_path: Path) -> None:
    """Write the complete accepted list without quality-review input columns."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "営業リスト"
    dark = "17324D"

    sheet.merge_cells("A1:J1")
    sheet["A1"] = f"営業リスト｜確認済み{len(records)}社"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=dark)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells("A2:J2")
    sheet["A2"] = "営業に使う項目だけを表示しています。詳しい判定根拠は audit.json にあります。"
    sheet["A2"].font = Font(color="425466")

    for column, header in enumerate(SALES_HEADERS, 1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[4].height = 28

    for row_number, record in enumerate(records, 5):
        for column, value in enumerate(sales_row(record, row_number - 4), 1):
            sheet.cell(row=row_number, column=column, value=value).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
        for column in (3, 7):
            cell = sheet.cell(row=row_number, column=column)
            if cell.value and " / " not in str(cell.value):
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    last_row = 4 + len(records)
    if records:
        _add_table(sheet, "VerifiedSalesLeads", f"A4:J{last_row}")
    widths = {
        "A": 7, "B": 28, "C": 34, "D": 20, "E": 40,
        "F": 28, "G": 36, "H": 42, "I": 42, "J": 22,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B5"
    sheet.auto_filter.ref = f"A4:J{last_row}"
    sheet.sheet_view.showGridLines = False
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_review_workbook(
    records: list[dict[str, Any]],
    output_path: Path,
    human_labels: dict[str, tuple[str, str]] | None = None,
) -> None:
    """Write a selected sample as a filterable quality-review workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "確認リスト"
    dark = "17324D"
    blue = "DDEBF7"
    green = "E2F0D9"
    red = "FCE4D6"

    sheet.merge_cells("A1:P1")
    sheet["A1"] = f"営業リスト｜抜き取り品質チェック（{len(records)}社）"
    sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=dark)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 32
    sheet["A2"] = "使い方"
    sheet["B2"] = "公式URLと3種類の証拠を見て、B列を「正しい」または「間違い」にしてください。気づいた点はC列へ書きます。"
    sheet.merge_cells("B2:P2")
    sheet["A2"].font = Font(bold=True, color=dark)
    sheet["B2"].alignment = Alignment(wrap_text=True)

    sheet["A4"] = "チェック対象"
    sheet["B4"] = len(records)
    sheet["D4"] = "確認済み"
    sheet["E4"] = '=COUNTIF(B7:B1048576,"正しい")+COUNTIF(B7:B1048576,"間違い")'
    sheet["G4"] = "正しい"
    sheet["H4"] = '=COUNTIF(B7:B1048576,"正しい")'
    sheet["J4"] = "間違い"
    sheet["K4"] = '=COUNTIF(B7:B1048576,"間違い")'
    for coordinate in ("A4", "D4", "G4", "J4"):
        sheet[coordinate].font = Font(bold=True, color=dark)
        sheet[coordinate].fill = PatternFill("solid", fgColor=blue)

    for column, header in enumerate(REVIEW_HEADERS, 1):
        cell = sheet.cell(row=6, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row_number, record in enumerate(records, 7):
        for column, value in enumerate(review_row(record, row_number - 6, human_labels), 1):
            sheet.cell(row=row_number, column=column, value=value).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
        for column in (5, 9, 11, 13, 15):
            cell = sheet.cell(row=row_number, column=column)
            if cell.value and " / " not in str(cell.value):
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    last_row = 6 + len(records)
    if records:
        _add_table(sheet, "VerifiedLeadsReview", f"A6:P{last_row}")
        verdicts = DataValidation(type="list", formula1='"正しい,間違い"', allow_blank=True)
        verdicts.promptTitle = "人の確認"
        verdicts.prompt = "正しい、または間違いを選んでください。"
        verdicts.error = "一覧から選んでください。"
        verdicts.errorTitle = "入力を確認してください"
        verdicts.showErrorMessage = True
        sheet.add_data_validation(verdicts)
        verdicts.add(f"B7:B{last_row}")
        sheet.conditional_formatting.add(
            f"A7:P{last_row}",
            FormulaRule(formula=['$B7="正しい"'], fill=PatternFill("solid", fgColor=green)),
        )
        sheet.conditional_formatting.add(
            f"A7:P{last_row}",
            FormulaRule(formula=['$B7="間違い"'], fill=PatternFill("solid", fgColor=red)),
        )

    widths = {
        "A": 7, "B": 12, "C": 24, "D": 28, "E": 34, "F": 18, "G": 34, "H": 26,
        "I": 34, "J": 44, "K": 34, "L": 44, "M": 34, "N": 44, "O": 34, "P": 22,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "D7"
    sheet.auto_filter.ref = f"A6:P{last_row}"
    sheet.sheet_view.showGridLines = False

    guide = workbook.create_sheet("確認ガイド")
    guide["A1"] = f"{len(records)}社の抜き取り品質チェック"
    guide["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    guide["A1"].fill = PatternFill("solid", fgColor=dark)
    guide.merge_cells("A1:F1")
    steps = [
        ("1", "公式URLを開く", "会社名とサイトが一致し、ページが表示されるか確認します。"),
        ("2", "商材を見る", "売りたい商材との関係が、証拠文とページで分かるか確認します。"),
        ("3", "会社の役割を見る", "カフェ、卸会社、輸入会社など、今回会いたい相手か確認します。"),
        ("4", "対象地域を見る", "対象国で店舗・事業所・販売活動があるか確認します。"),
        ("5", "B列で判定する", "3つとも問題なければ「正しい」。対象外や誤りなら「間違い」にします。"),
    ]
    for row, (number, title, detail) in enumerate(steps, 3):
        guide.cell(row, 1, number).font = Font(bold=True, color="FFFFFF")
        guide.cell(row, 1).fill = PatternFill("solid", fgColor=dark)
        guide.cell(row, 2, title).font = Font(bold=True, color=dark)
        guide.cell(row, 3, detail).alignment = Alignment(wrap_text=True)
        guide.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        guide.row_dimensions[row].height = 42
    guide["A10"] = "結果の目安"
    guide["A10"].font = Font(bold=True, color=dark)
    guide["A11"] = "9〜10社が正しい"
    guide["B11"] = "簡易チェックは合格です。これは全件の精度を断定するものではありません。"
    guide["A12"] = "8社以下が正しい"
    guide["B12"] = "20社または全件へ確認を広げ、間違いの原因を調べます。"
    guide["A11"].fill = PatternFill("solid", fgColor=green)
    guide["A12"].fill = PatternFill("solid", fgColor=red)
    guide.merge_cells("B11:F11")
    guide.merge_cells("B12:F12")
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 28
    for column in ("C", "D", "E", "F"):
        guide.column_dimensions[column].width = 18
    guide.sheet_view.showGridLines = False
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--mode", choices=("sales", "review"), default="review")
    parser.add_argument("--sample-size", type=int, default=10)
    parser.add_argument("--all", action="store_true", help="include every accepted record in review mode")
    args = parser.parse_args()
    records = json.loads(args.audit.read_text(encoding="utf-8"))
    accepted = [item for item in records if item.get("verification_status") == "accepted"]
    if args.mode == "sales":
        exported = accepted
        write_sales_workbook(exported, args.output)
        sheet_name = "営業リスト"
        expected_rows = 4 + len(exported)
    else:
        sample_size = len(accepted) if args.all else args.sample_size
        exported = select_quality_sample(accepted, sample_size)
        write_review_workbook(exported, args.output)
        sheet_name = "確認リスト"
        expected_rows = 6 + len(exported)
    loaded = load_workbook(args.output, read_only=True)
    if loaded[sheet_name].max_row != expected_rows:
        raise RuntimeError("Excel row count verification failed")
    print(json.dumps({"output": str(args.output), "mode": args.mode, "rows": len(exported)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
